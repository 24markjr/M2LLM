"""Generate the Orion fixtures (Phase 41): the evaluation's first Word, Excel, image and subtitle files.

Run once from the repository root; the outputs are committed, so evaluation never depends on this
script or on the fonts of the machine that runs it:

    python .agent/fixtures/make_orion_fixtures.py

A fourth document family after Aurora, Helix and the shipment, built to make the scenarios read
through the Phase 38-39 parsers rather than plain text:

- `orion_budget_memo.docx` + `orion_ledger.xlsx`: the memo approves INR 250,000; the ledger's total
  is INR 287,500. The overrun is the planted finding (scenario `orion_ledger_overrun`).
- `orion_purchase_order.txt` + `orion_delivery_note.png`: 40 units ordered, a scanned note says 36
  delivered. The note is an image, so the finding exists only if OCR read it
  (`orion_scanned_delivery`).
- `orion_meeting.srt` + `orion_minutes.docx`: the recording and the minutes agree on the handover
  date. Nothing to find (`orion_meeting_consistent`, a negative case).
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont

OUT = Path(__file__).resolve().parent / "documents"


def memo() -> None:
    document = Document()
    document.add_heading("PROJECT ORION - BUDGET APPROVAL MEMO", level=1)
    document.add_paragraph("From: Steering Committee, Orion Cold Storage Programme")
    document.add_paragraph("Date: 3 March 2026")
    document.add_paragraph(
        "The steering committee approved a total budget of INR 250,000 for Project Orion, "
        "covering equipment, installation and commissioning."
    )
    document.add_paragraph(
        "Any spend above the approved budget requires a change request signed by the programme "
        "sponsor before it is incurred."
    )
    document.add_paragraph("No change request has been raised as of the date of this memo.")
    document.save(OUT / "orion_budget_memo.docx")


def ledger() -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Spend"
    rows = [
        ("Item", "Supplier", "Date", "Amount (INR)"),
        ("Cooling units (36)", "Polar Systems", "2026-03-19", 162000),
        ("Installation", "Polar Systems", "2026-03-24", 58500),
        ("Electrical works", "Sundar Electricals", "2026-03-26", 41000),
        ("Commissioning", "Polar Systems", "2026-04-01", 26000),
        ("Total spend", "", "", 287500),
    ]
    for row in rows:
        sheet.append(row)
    notes = workbook.create_sheet("Notes")
    notes.append(("Ledger maintained by the Orion finance office.",))
    notes.append(("Figures are actual invoiced amounts.",))
    workbook.save(OUT / "orion_ledger.xlsx")


def purchase_order() -> None:
    (OUT / "orion_purchase_order.txt").write_text(
        "PROJECT ORION - PURCHASE ORDER PO-7741\n"
        "\n"
        "Supplier: Polar Systems\n"
        "Issued: 5 March 2026\n"
        "\n"
        "Item: Cooling units, model PX-200\n"
        "Quantity ordered: 40 units\n"
        "Delivery due: 12 March 2026, Orion site, Pune\n"
        "\n"
        "Partial deliveries must be agreed in writing by the Orion procurement office.\n",
        encoding="utf-8",
    )


def delivery_note() -> None:
    lines = [
        "POLAR SYSTEMS - DELIVERY NOTE DN-3310",
        "Order reference: PO-7741",
        "Delivered to: Orion site, Pune",
        "Delivery date: 19 March 2026",
        "Item: Cooling units, model PX-200",
        "Units delivered: 36",
        "Received by: R. Kulkarni",
    ]
    try:
        font = ImageFont.truetype("arial.ttf", 34)
    except OSError:
        font = ImageFont.truetype("DejaVuSans.ttf", 34)
    image = Image.new("RGB", (1100, 70 + 60 * len(lines)), "white")
    draw = ImageDraw.Draw(image)
    for number, line in enumerate(lines):
        draw.text((40, 35 + 60 * number), line, fill="black", font=font)
    image.save(OUT / "orion_delivery_note.png")


def meeting() -> None:
    cues = [
        ("00:00:01,000", "00:00:05,000", "Good morning. This is the Orion site readiness review."),
        ("00:00:05,500", "00:00:11,000", "Installation finished on 1 April and commissioning passed."),
        ("00:00:11,500", "00:00:17,000", "So we confirm the site handover to operations on 6 April 2026."),
        ("00:00:17,500", "00:00:22,000", "Operations will take over the night shift from that date."),
    ]
    blocks = [f"{n}\n{start} --> {end}\n{text}\n" for n, (start, end, text) in enumerate(cues, 1)]
    (OUT / "orion_meeting.srt").write_text("\n".join(blocks), encoding="utf-8")

    document = Document()
    document.add_heading("PROJECT ORION - SITE READINESS REVIEW, MINUTES", level=1)
    document.add_paragraph("Attendees: programme sponsor, site lead, operations lead")
    document.add_paragraph("1. Installation completed on 1 April 2026; commissioning passed.")
    document.add_paragraph("2. Site handover to operations confirmed for 6 April 2026.")
    document.add_paragraph("3. Operations assumes the night shift from the handover date.")
    document.save(OUT / "orion_minutes.docx")


if __name__ == "__main__":
    memo()
    ledger()
    purchase_order()
    delivery_note()
    meeting()
    print(f"wrote the Orion fixtures to {OUT}")
