"""Generate the Kestrel fixtures (Phase B2): the multimodal evaluation's audio, video, scan and photo.

Run once from the repository root on Windows (the speech is Windows' own synthesiser); the outputs are
committed, so evaluation never depends on this script, its fonts or its voice:

    python .agent/fixtures/make_kestrel_fixtures.py

One document family, one planted finding per kind of file, each existing only inside the non-text
file - so it is found only if the file was understood:

- `kestrel_call.wav` (speech) + `kestrel_minutes.docx`: the call agrees go-live on 9 June 2026; the
  minutes record 16 June 2026. Both say the budget cap is INR 120,000 (the negative case).
- `kestrel_site.mp4` (on-screen text, no sound) + `kestrel_inspection_report.txt`: the video shows
  generator B failing its load test on 2 June 2026; the report says every generator passed.
- `kestrel_invoice_scan.png` + `kestrel_licences.xlsx`: the scanned invoice bills 48 licences for
  INR 96,000; the order sheet approved 40 licences for INR 80,000.
- `kestrel_whiteboard.jpg` (a skewed, noisy photo) + `kestrel_restore.log`: the whiteboard says the
  backup restore finished at 03:40; the log records it completing at 05:15.
"""

from __future__ import annotations

import random
import subprocess
from pathlib import Path

import av
import numpy as np
from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFilter, ImageFont

OUT = Path(__file__).resolve().parent / "documents"


def font(size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    for name in ("arial.ttf", "DejaVuSans.ttf"):
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            continue
    return ImageFont.load_default()


def text_image(lines: list[str], size: tuple[int, int], *, title_size: int = 34, body_size: int = 28) -> Image.Image:
    image = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(image)
    y = 40
    for i, line in enumerate(lines):
        draw.text((50, y), line, fill="black", font=font(title_size if i == 0 else body_size))
        y += (title_size if i == 0 else body_size) + 22
    return image


def audio() -> None:
    speech = (
        "This is the Kestrel steering call. We agreed today that the go-live date is the ninth of June "
        "twenty twenty six. The budget cap stays at one hundred and twenty thousand rupees. "
        "Anika will circulate the minutes."
    )
    path = OUT / "kestrel_call.wav"
    script = (
        "Add-Type -AssemblyName System.Speech; "
        "$s = New-Object System.Speech.Synthesis.SpeechSynthesizer; "
        f"$s.SetOutputToWaveFile('{path}'); $s.Speak('{speech}'); $s.Dispose()"
    )
    subprocess.run(["powershell", "-NoProfile", "-Command", script], check=True)  # noqa: S603, S607

    doc = Document()
    doc.add_heading("Kestrel steering committee - minutes", level=1)
    doc.add_paragraph("Attendees: Anika Rao (chair), Dev Malhotra, Sara Iqbal.")
    doc.add_paragraph("Decision 1: The go-live date is 16 June 2026.")
    doc.add_paragraph("Decision 2: The budget cap remains INR 120,000.")
    doc.add_paragraph("Minutes circulated by Anika Rao.")
    doc.save(str(OUT / "kestrel_minutes.docx"))


def video() -> None:
    frames = [
        ["KESTREL SITE INSPECTION", "Date: 2 June 2026", "Inspector: Dev Malhotra"],
        ["GENERATOR A - LOAD TEST", "Result: PASSED", "Load held at 100% for 30 minutes"],
        ["GENERATOR B - LOAD TEST", "Result: FAILED", "Shut down at 70% load after 4 minutes"],
        ["ACTION", "Generator B to be repaired and retested", "before go-live"],
    ]
    path = OUT / "kestrel_site.mp4"
    with av.open(str(path), mode="w") as container:
        stream = container.add_stream("mpeg4", rate=1)
        stream.width, stream.height, stream.pix_fmt = 960, 540, "yuv420p"
        for lines in frames:
            picture = np.array(text_image(lines, (960, 540), title_size=40, body_size=34))
            for _ in range(4):  # four seconds per card
                for packet in stream.encode(av.VideoFrame.from_ndarray(picture, format="rgb24")):
                    container.mux(packet)
        for packet in stream.encode():
            container.mux(packet)

    (OUT / "kestrel_inspection_report.txt").write_text(
        "KESTREL - SITE INSPECTION REPORT\n"
        "Report ID: KES-INS-2026-06\n"
        "Inspection date: 2 June 2026\n"
        "\n"
        "1. GENERATORS\n"
        "All generators passed the load test on 2 June 2026.\n"
        "\n"
        "2. READINESS\n"
        "The site is ready for go-live.\n",
        encoding="utf-8",
    )


def scan() -> None:
    image = text_image(
        [
            "TAX INVOICE  INV-KES-0412",
            "Supplier: Northwind Software",
            "Customer: Kestrel Programme",
            "Date: 28 May 2026",
            "Item: Analytics licence x 48",
            "Unit price: INR 2,000",
            "Total: INR 96,000",
        ],
        (1000, 640),
    )
    # A scan: slightly rotated, softened, grey speckle.
    image = image.rotate(1.2, expand=True, fillcolor="white").filter(ImageFilter.GaussianBlur(0.6))
    rng = random.Random(7)
    draw = ImageDraw.Draw(image)
    for _ in range(1800):
        x, y = rng.randrange(image.width), rng.randrange(image.height)
        draw.point((x, y), fill=(rng.randrange(150, 220),) * 3)
    image.save(OUT / "kestrel_invoice_scan.png")

    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Approved order"
    for row in [
        ("Item", "Quantity", "Unit price (INR)", "Total (INR)", "Approved by"),
        ("Analytics licence", 40, 2000, 80000, "Sara Iqbal"),
    ]:
        sheet.append(row)
    book.save(str(OUT / "kestrel_licences.xlsx"))


def photo() -> None:
    image = Image.new("RGB", (1100, 760), (236, 236, 228))
    draw = ImageDraw.Draw(image)
    draw.rectangle((60, 50, 1040, 700), fill=(250, 250, 246), outline=(120, 120, 120), width=6)
    y = 110
    for line, size in [
        ("DB RESTORE - NIGHT OF 3 JUNE", 46),
        ("Started 01:05", 40),
        ("Restore finished 03:40", 40),
        ("Checked by Sara", 40),
    ]:
        draw.text((110, y), line, fill=(20, 40, 120), font=font(size))
        y += size + 60
    # A photo, not a scan: perspective-ish skew, uneven light, blur, sensor noise.
    image = image.rotate(-3.5, expand=True, fillcolor=(200, 200, 195)).filter(ImageFilter.GaussianBlur(1.1))
    pixels = np.array(image).astype(np.int16)
    shade = np.linspace(-25, 20, pixels.shape[1], dtype=np.int16)[None, :, None]
    noise = np.random.default_rng(11).normal(0, 7, pixels.shape).astype(np.int16)
    Image.fromarray(np.clip(pixels + shade + noise, 0, 255).astype(np.uint8)).save(
        OUT / "kestrel_whiteboard.jpg", quality=82
    )

    (OUT / "kestrel_restore.log").write_text(
        "2026-06-04T01:05:12Z restore-job INFO restore started source=backup-2026-06-03\n"
        "2026-06-04T02:30:40Z restore-job INFO 61% of tables restored\n"
        "2026-06-04T05:15:03Z restore-job INFO restore completed tables=412 errors=0\n"
        "2026-06-04T05:16:10Z restore-job INFO integrity check passed\n",
        encoding="utf-8",
    )


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    audio()
    video()
    scan()
    photo()
    print("written to", OUT)  # noqa: T201
