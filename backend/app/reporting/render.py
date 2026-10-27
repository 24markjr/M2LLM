"""The report document drawn three ways: HTML, PDF and DOCX (v4 Phase 5).

All three draw `build_document(report)` block by block, so screen, PDF and DOCX hold the same words
in the same order; only the drawing differs. The tone of a cell is a colour and a word - the word is
always printed, so a black-and-white printout loses nothing.

- **HTML**: one self-contained page (inline CSS, no scripts), for the screen and for printing.
- **PDF**: ReportLab's layout engine (BSD), A4, page numbers and the report's hashes on every page.
  A Unicode font (DejaVu Sans, or Arial on Windows) is used when the machine has one, so the rupee
  sign prints; otherwise Helvetica, with such characters spelled out ("INR").
- **DOCX**: python-docx, real tables with shaded status cells, so a reader can edit and re-use it.
"""

from __future__ import annotations

import html
import io
from functools import lru_cache
from pathlib import Path

from app.reporting.document import Block, Cell, ReportDocument, Tone
from app.reporting.export import _ASCII

# Background and text colour of each tone, shared by the three renderers.
TONES: dict[str, tuple[str, str]] = {
    "verified": ("#E3F4EA", "#1D6B3E"),
    "partial": ("#FFF3D6", "#8A5A00"),
    "unsupported": ("#FBE3E3", "#9B1C1C"),
    "critical": ("#7F1D1D", "#FFFFFF"),
    "high": ("#FBE3E3", "#9B1C1C"),
    "medium": ("#FFF3D6", "#8A5A00"),
    "low": ("#EAF0FB", "#28447A"),
    "muted": ("#F1F3F6", "#5A6372"),
}
INK, ACCENT, RULE = "#1B2230", "#2C56E0", "#D9DEE6"


def _tone(tone: Tone) -> tuple[str, str] | None:
    return TONES.get(tone) if tone else None


# --- HTML ----------------------------------------------------------------------------------------

_CSS = f"""
* {{ box-sizing: border-box; }}
body {{ font: 14px/1.55 "Segoe UI", Arial, "DejaVu Sans", sans-serif; color: {INK}; margin: 0;
  background: #F6F7F9; }}
main {{ max-width: 960px; margin: 24px auto; background: #fff; padding: 40px 48px;
  border: 1px solid {RULE}; border-radius: 10px; }}
header.cover {{ border-bottom: 3px solid {ACCENT}; padding-bottom: 18px; margin-bottom: 8px; }}
h1 {{ font-size: 28px; margin: 0 0 6px; }}
.subtitle {{ font-size: 16px; color: #3A4456; margin: 0 0 14px; }}
h2 {{ font-size: 18px; margin: 30px 0 10px; padding-bottom: 6px; border-bottom: 1px solid {RULE}; }}
dl.facts {{ display: grid; grid-template-columns: max-content 1fr; gap: 4px 16px; margin: 10px 0; }}
dl.facts dt {{ color: #5A6372; }} dl.facts dd {{ margin: 0; overflow-wrap: anywhere; }}
table {{ width: 100%; border-collapse: collapse; margin: 8px 0 4px; font-size: 13px; }}
th {{ text-align: left; background: #F1F3F6; padding: 7px 8px; border-bottom: 2px solid {RULE}; }}
td {{ padding: 7px 8px; border-bottom: 1px solid {RULE}; vertical-align: top;
  overflow-wrap: anywhere; }}
.tone {{ display: inline-block; padding: 1px 8px; border-radius: 10px; font-weight: 600;
  font-size: 12px; white-space: nowrap; }}
p.lead {{ padding: 10px 14px; border-radius: 8px; font-weight: 600; }}
.pair {{ margin: 12px 0; border: 1px solid {RULE}; border-radius: 8px; overflow: hidden; }}
.pair .claim {{ padding: 8px 12px; background: #F1F3F6; font-weight: 600; }}
.pair .sides {{ display: grid; grid-template-columns: 1fr 1fr; }}
.pair .side {{ padding: 10px 12px; }} .pair .side + .side {{ border-left: 1px solid {RULE}; }}
.pair .src {{ font: 12px ui-monospace, Consolas, monospace; color: #5A6372; }}
.pair blockquote {{ margin: 6px 0 0; font-style: italic; }}
footer {{ margin-top: 32px; color: #5A6372; font-size: 12px; border-top: 1px solid {RULE};
  padding-top: 10px; overflow-wrap: anywhere; }}
@media (max-width: 640px) {{ main {{ padding: 20px 16px; margin: 0; border-radius: 0; }}
  .pair .sides {{ grid-template-columns: 1fr; }} table {{ display: block; overflow-x: auto; }} }}
@media print {{ body {{ background: #fff; }} main {{ border: 0; margin: 0; padding: 0; }}
  h2 {{ break-after: avoid; }} tr, .pair {{ break-inside: avoid; }} }}
"""


def _cell_html(cell: Cell) -> str:
    text = "<br>".join(html.escape(part) for part in cell.text.split("\n"))
    colours = _tone(cell.tone)
    if colours is None:
        return text
    return f'<span class="tone" style="background:{colours[0]};color:{colours[1]}">{text}</span>'


def _block_html(block: Block) -> str:
    e = html.escape
    if block.kind == "heading":
        return f"<h2>{e(block.text)}</h2>"
    if block.kind == "paragraph":
        colours = _tone(block.tone)
        if colours is None:
            return f"<p>{e(block.text)}</p>"
        style = f"background:{colours[0]};color:{colours[1]}"
        return f'<p class="lead" style="{style}">{e(block.text)}</p>'
    if block.kind == "bullets":
        return "<ul>" + "".join(f"<li>{e(item)}</li>" for item in block.items) + "</ul>"
    if block.kind == "facts":
        return (
            '<dl class="facts">'
            + "".join(f"<dt>{e(k)}</dt><dd>{e(v)}</dd>" for k, v in block.facts)
            + "</dl>"
        )
    if block.kind == "pair":
        sides = "".join(
            f'<div class="side"><div class="src">{e(src)}</div>'
            f'<blockquote>"{e(quote)}"</blockquote></div>'
            for src, quote in (block.left, block.right)
        )
        return (
            f'<div class="pair"><div class="claim">{e(block.text)}</div>'
            f'<div class="sides">{sides}</div></div>'
        )
    head = "".join(f"<th>{e(c)}</th>" for c in block.columns)
    rows = "".join(
        "<tr>" + "".join(f"<td>{_cell_html(c)}</td>" for c in row) + "</tr>" for row in block.rows
    )
    return f"<table><thead><tr>{head}</tr></thead><tbody>{rows}</tbody></table>"


def render_html(document: ReportDocument) -> str:
    e = html.escape
    facts = "".join(f"<dt>{e(k)}</dt><dd>{e(v)}</dd>" for k, v in document.facts)
    body = "\n".join(_block_html(b) for b in document.blocks)
    return (
        '<!doctype html>\n<html lang="en"><head><meta charset="utf-8">'
        '<meta name="viewport" content="width=device-width, initial-scale=1">'
        f"<title>{e(document.title)}</title><style>{_CSS}</style></head><body><main>"
        f'<header class="cover"><h1>{e(document.title)}</h1>'
        f'<p class="subtitle">{e(document.subtitle)}</p><dl class="facts">{facts}</dl></header>'
        f"{body}<footer>{e(document.footer)}</footer></main></body></html>\n"
    )


# --- PDF -----------------------------------------------------------------------------------------

_FONT_CANDIDATES = [
    (
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
    ),
    ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
    ("/Library/Fonts/Arial Unicode.ttf", "/Library/Fonts/Arial Unicode.ttf"),
]


def _column_widths(block: Block, width: float, font: str) -> list[float]:
    """Each column as wide as its longest word (so no word is split), the widest-text column taking
    what is left. A column whose words cannot all fit shares the shortfall in proportion."""
    from reportlab.pdfbase.pdfmetrics import stringWidth

    padding = 12.0
    columns = len(block.columns)
    words = [
        [block.columns[i], *(w for row in block.rows for w in row[i].text.split())]
        for i in range(columns)
    ]
    natural = [max(stringWidth(w, font, 8.5) for w in column) + padding for column in words]
    lengths = [sum(len(row[i].text) for row in block.rows) for i in range(columns)]
    widest = lengths.index(max(lengths)) if lengths else 0
    others = sum(n for i, n in enumerate(natural) if i != widest)
    room = width * 0.62  # the widest column keeps at least 38 % of the line
    scale = min(1.0, room / others) if others else 1.0
    widths = [n * scale for n in natural]
    widths[widest] = width - sum(w for i, w in enumerate(widths) if i != widest)
    return widths


@lru_cache(maxsize=1)
def _fonts() -> tuple[str, str, bool]:
    """(regular, bold, unicode): a TrueType font when one is installed, else built-in Helvetica."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    for regular, bold in _FONT_CANDIDATES:
        if Path(regular).is_file() and Path(bold).is_file():
            pdfmetrics.registerFont(TTFont("Report", regular))
            pdfmetrics.registerFont(TTFont("Report-Bold", bold))
            return "Report", "Report-Bold", True
    return "Helvetica", "Helvetica-Bold", False


def render_pdf(document: ReportDocument) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        KeepTogether,
        ListFlowable,
        ListItem,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    regular, bold, unicode = _fonts()

    def text(value: str) -> str:
        value = (
            value
            if unicode
            else value.translate(_ASCII).encode("latin-1", "replace").decode("latin-1")
        )
        return html.escape(value)

    def lines(value: str) -> str:
        """A cell's text, its line breaks kept (a status is its standing over its verdict)."""
        return "<br/>".join(text(part) for part in value.split("\n"))

    body = ParagraphStyle(
        "body", fontName=regular, fontSize=9.5, leading=13, textColor=colors.HexColor(INK)
    )
    small = ParagraphStyle("small", parent=body, fontSize=8.5, leading=11)
    title = ParagraphStyle(
        "title", parent=body, fontName=bold, fontSize=22, leading=27, spaceAfter=4
    )
    subtitle = ParagraphStyle(
        "subtitle", parent=body, fontSize=12, leading=16, textColor=colors.HexColor("#3A4456")
    )
    heading = ParagraphStyle(
        "heading",
        parent=body,
        fontName=bold,
        fontSize=13.5,
        leading=18,
        spaceBefore=14,
        spaceAfter=6,
        keepWithNext=1,  # a heading never ends a page alone
    )
    width = A4[0] - 36 * mm

    def facts_table(facts: list[tuple[str, str]]) -> Table:
        table = Table(
            [[Paragraph(text(k), small), Paragraph(text(v), body)] for k, v in facts],
            colWidths=[38 * mm, width - 38 * mm],
        )
        table.setStyle(
            TableStyle(
                [
                    ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor("#5A6372")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                ]
            )
        )
        return table

    story: list[object] = [
        Paragraph(text(document.title), title),
        Paragraph(text(document.subtitle), subtitle),
        Spacer(1, 4),
        facts_table(document.facts),
        Spacer(1, 4),
    ]
    for block in document.blocks:
        if block.kind == "heading":
            story.append(Paragraph(text(block.text), heading))
        elif block.kind == "paragraph":
            colours = _tone(block.tone)
            if colours is None:
                story.append(Paragraph(text(block.text), body))
            else:
                lead = Table([[Paragraph(f"<b>{text(block.text)}</b>", body)]], colWidths=[width])
                lead.setStyle(
                    TableStyle(
                        [
                            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(colours[0])),
                            ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor(colours[1])),
                            ("LEFTPADDING", (0, 0), (-1, -1), 8),
                            ("TOPPADDING", (0, 0), (-1, -1), 6),
                            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                        ]
                    )
                )
                story.append(lead)
        elif block.kind == "bullets":
            story.append(
                ListFlowable(
                    [ListItem(Paragraph(text(i), body), leftIndent=10) for i in block.items],
                    bulletType="bullet",
                    start="•" if unicode else "-",
                    bulletFontName=regular,
                    leftIndent=10,
                )
            )
        elif block.kind == "facts":
            story.append(facts_table(block.facts))
        elif block.kind == "pair":
            sides = [
                Paragraph(
                    f'<font size="8" color="#5A6372">{text(src)}</font><br/><i>"{text(q)}"</i>',
                    body,
                )
                for src, q in (block.left, block.right)
            ]
            pair = Table(
                [[Paragraph(f"<b>{text(block.text)}</b>", body), ""], sides],
                colWidths=[width / 2, width / 2],
            )
            pair.setStyle(
                TableStyle(
                    [
                        ("SPAN", (0, 0), (1, 0)),
                        ("BACKGROUND", (0, 0), (1, 0), colors.HexColor("#F1F3F6")),
                        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor(RULE)),
                        ("LINEBEFORE", (1, 1), (1, 1), 0.5, colors.HexColor(RULE)),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("TOPPADDING", (0, 0), (-1, -1), 5),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                    ]
                )
            )
            story.append(KeepTogether([pair, Spacer(1, 6)]))
        else:  # table
            widths = _column_widths(block, width, bold)
            data = [[Paragraph(f"<b>{text(c)}</b>", small) for c in block.columns]]
            data += [[Paragraph(lines(c.text), small) for c in row] for row in block.rows]
            style = [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F1F3F6")),
                ("LINEBELOW", (0, 0), (-1, -1), 0.4, colors.HexColor(RULE)),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
            for r, row in enumerate(block.rows, start=1):
                for c, cell in enumerate(row):
                    colours = _tone(cell.tone)
                    if colours is not None:
                        style.append(("BACKGROUND", (c, r), (c, r), colors.HexColor(colours[0])))
                        data[r][c] = Paragraph(
                            f'<font color="{colours[1]}"><b>{lines(cell.text)}</b></font>', small
                        )
            table = Table(data, colWidths=widths, repeatRows=1)
            table.setStyle(TableStyle(style))
            story.append(table)

    footer = text(document.footer)

    def decorate(canvas: object, doc: object) -> None:
        canvas.saveState()  # type: ignore[attr-defined]
        canvas.setFont(regular, 7.5)  # type: ignore[attr-defined]
        canvas.setFillColor(colors.HexColor("#5A6372"))  # type: ignore[attr-defined]
        canvas.drawString(18 * mm, 10 * mm, html.unescape(footer))  # type: ignore[attr-defined]
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Page {doc.page}")  # type: ignore[attr-defined]
        canvas.restoreState()  # type: ignore[attr-defined]

    buffer = io.BytesIO()
    SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=18 * mm,
        title=document.title,
        subject=document.footer,
        author="JARVIS",
    ).build(story, onFirstPage=decorate, onLaterPages=decorate)
    return buffer.getvalue()


# --- DOCX ----------------------------------------------------------------------------------------


def render_docx(document: ReportDocument) -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    def shade(cell: object, colour: str) -> None:
        properties = cell._tc.get_or_add_tcPr()  # type: ignore[attr-defined]
        fill = OxmlElement("w:shd")
        fill.set(qn("w:val"), "clear")
        fill.set(qn("w:color"), "auto")
        fill.set(qn("w:fill"), colour.lstrip("#"))
        properties.append(fill)

    def write(cell: object, value: str, tone: Tone = "", bold: bool = False) -> None:
        paragraph = cell.paragraphs[0]  # type: ignore[attr-defined]
        run = paragraph.add_run(value)
        run.font.size = Pt(8.5)
        run.bold = bold
        colours = _tone(tone)
        if colours is not None:
            shade(cell, colours[0])
            run.font.color.rgb = RGBColor.from_string(colours[1].lstrip("#"))
            run.bold = True

    def facts(rows: list[tuple[str, str]]) -> None:
        table = out.add_table(rows=0, cols=2)
        for key, value in rows:
            cells = table.add_row().cells
            write(cells[0], key)
            write(cells[1], value)

    out = Document()
    style = out.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)
    out.add_heading(document.title, level=0)
    subtitle = out.add_paragraph(document.subtitle)
    subtitle.runs[0].font.size = Pt(12)
    facts(document.facts)
    for block in document.blocks:
        if block.kind == "heading":
            out.add_heading(block.text, level=1)
        elif block.kind == "paragraph":
            paragraph = out.add_paragraph()
            run = paragraph.add_run(block.text)
            colours = _tone(block.tone)
            if colours is not None:
                run.bold = True
                run.font.color.rgb = RGBColor.from_string(colours[1].lstrip("#"))
        elif block.kind == "bullets":
            for item in block.items:
                out.add_paragraph(item, style="List Bullet")
        elif block.kind == "facts":
            facts(block.facts)
        elif block.kind == "pair":
            table = out.add_table(rows=2, cols=2)
            table.style = "Table Grid"
            merged = table.cell(0, 0).merge(table.cell(0, 1))
            write(merged, block.text, bold=True)
            shade(merged, "#F1F3F6")
            for column, (src, quote) in enumerate((block.left, block.right)):
                write(table.cell(1, column), f'{src}\n"{quote}"')
            out.add_paragraph()
        else:
            table = out.add_table(rows=1, cols=len(block.columns))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for cell, name in zip(table.rows[0].cells, block.columns, strict=True):
                write(cell, name, bold=True)
                shade(cell, "#F1F3F6")
            for row in block.rows:
                cells = table.add_row().cells
                for cell, value in zip(cells, row, strict=True):
                    write(cell, value.text, value.tone)
    section = out.sections[0]
    section.footer.paragraphs[0].text = document.footer
    buffer = io.BytesIO()
    out.save(buffer)
    return buffer.getvalue()
