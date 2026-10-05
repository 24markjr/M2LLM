"""Every file type a mission can read, and the parser for each (Phase 38, Member 2's formats).

Member 2's M2LLM classified a file by its extension and routed it to one parser per type. So does
this registry, and it is the single source of truth: the upload endpoint accepts exactly these
extensions, the console's file picker offers exactly these (`GET /documents/formats`), and
`load_document` dispatches through them. A type cannot be accepted somewhere and unreadable
somewhere else.

**Everything becomes lines.** Whatever the format, a document is turned into numbered lines, so a
finding cites `file:rN` and the verifier reads that line - the same evidence path for a Word table
as for a text file. Formats with natural divisions keep them as "pages", the mechanism PDFs already
use: a spreadsheet's sheets are its pages, so a citation can say which sheet.

**A file with no text says so.** An image or a video is accepted and described (size, format,
orientation; the sidecar transcript if there is one), and marked `has_text=False` until text is
extracted from it: OCR for images and scans arrives in Phase 39. A mission excludes such a file and
logs why, rather than treating a silent file as evidence that said nothing.
"""

from __future__ import annotations

import hashlib
import re
from collections.abc import Callable
from enum import StrEnum
from pathlib import Path

from app.schemas.common import JarvisModel
from app.tools.loader import (
    MAX_DOCUMENT_CHARS,
    TRUNCATION_MARKER,
    DocumentLoadError,
    LoadedDocument,
    load_pdf,
    load_text,
)


class FormatKind(StrEnum):
    TEXT = "text"
    TABLE = "csv"  # the value the loader always used for CSV, kept so stored runs read the same
    SPREADSHEET = "spreadsheet"
    WORD = "word"
    PDF = "pdf"
    SUBTITLES = "subtitles"
    IMAGE = "image"
    VIDEO = "video"


class FormatSpec(JarvisModel):
    """One family of file types: which extensions, which parser, and how it becomes text."""

    kind: FormatKind
    extensions: list[str]
    parser: str
    # Whether a file of this kind yields citable text today. False for images and video until
    # OCR and transcription (Phase 39); a sidecar transcript can still give a video text.
    yields_text: bool = True
    becomes: str


FORMATS: list[FormatSpec] = [
    FormatSpec(
        kind=FormatKind.TEXT,
        extensions=[".txt", ".md", ".json", ".log"],
        parser="text",
        becomes="its lines, as written",
    ),
    FormatSpec(
        kind=FormatKind.TABLE,
        extensions=[".csv", ".tsv"],
        parser="text",
        becomes="one row per line",
    ),
    FormatSpec(
        kind=FormatKind.SPREADSHEET,
        extensions=[".xlsx", ".xlsm"],
        parser="openpyxl",
        becomes="one row per line, each sheet a page named in a marker line",
    ),
    FormatSpec(
        kind=FormatKind.WORD,
        extensions=[".docx"],
        parser="python-docx",
        becomes="paragraphs and table rows, in document order, one per line",
    ),
    FormatSpec(
        kind=FormatKind.PDF,
        extensions=[".pdf"],
        parser="pypdf",
        becomes="the text layer, page by page (a scanned page has none until OCR)",
    ),
    FormatSpec(
        kind=FormatKind.SUBTITLES,
        extensions=[".srt", ".vtt"],
        parser="subtitles",
        becomes="one cue per line, with its start time",
    ),
    FormatSpec(
        kind=FormatKind.IMAGE,
        extensions=[".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff", ".webp"],
        parser="pillow",
        yields_text=False,
        becomes="its size, format and orientation; text needs OCR (Phase 39)",
    ),
    FormatSpec(
        kind=FormatKind.VIDEO,
        extensions=[".mp4", ".mov", ".m4v", ".mkv", ".webm", ".avi"],
        parser="video",
        yields_text=False,
        becomes="its transcript when a subtitle file of the same name sits beside it",
    ),
]

_BY_EXTENSION: dict[str, FormatSpec] = {ext: spec for spec in FORMATS for ext in spec.extensions}

# A spreadsheet is read the way a CSV is: bounded, and said to be when it is cut.
MAX_SHEET_ROWS = 2000
MAX_SHEETS = 20


def spec_for(path: Path) -> FormatSpec | None:
    return _BY_EXTENSION.get(path.suffix.lower())


def supported_extensions() -> set[str]:
    return set(_BY_EXTENSION)


def sha256_of(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1 << 16), b""):
            digest.update(block)
    return digest.hexdigest()


def parse(path: Path) -> LoadedDocument:
    """Read one file with the parser its extension calls for, and record the provenance."""
    if not path.exists():
        raise DocumentLoadError(f"no such file: {path}")
    spec = spec_for(path)
    if spec is None:
        supported = ", ".join(sorted(supported_extensions()))
        raise DocumentLoadError(f"{path.name}: unsupported type; supported types are {supported}")
    document = _PARSERS[spec.kind](path)
    return document.model_copy(
        update={"kind": spec.kind.value, "parser": spec.parser, "sha256": sha256_of(path)}
    )


# --- the parsers ------------------------------------------------------------------------------


def _bounded(
    path: Path,
    lines: list[str],
    *,
    kind: str,
    page_starts: list[int] | None = None,
    page_count: int = 0,
    note: str = "",
) -> LoadedDocument:
    """Lines to a document, cut at the character ceiling every format shares."""
    text = "\n".join(lines)
    truncated = len(text) > MAX_DOCUMENT_CHARS
    if truncated:
        text = text[:MAX_DOCUMENT_CHARS] + f"\n{TRUNCATION_MARKER}"
        note = note or f"text beyond {MAX_DOCUMENT_CHARS} characters was not read"
    return LoadedDocument(
        document_id=path.name,
        text=text,
        kind=kind,
        truncated=truncated or bool(note),
        omitted_note=note,
        page_count=page_count,
        line_count=text.count("\n") + 1,
        size_bytes=path.stat().st_size,
        page_starts=page_starts or [],
    )


def load_docx(path: Path) -> LoadedDocument:
    """Paragraphs and tables, in the order they appear.

    Member 2's parser read paragraphs only, so a table - often where a report's figures are - was
    lost. Here each table row is one line, cells joined by " | ", after a marker naming the table.
    """
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise DocumentLoadError(f"python-docx is required to read {path.name}") from exc

    try:
        document = Document(str(path))
        lines: list[str] = []
        tables = 0
        for block in document.iter_inner_content():
            if isinstance(block, Paragraph):
                text = block.text.strip()
                if text:
                    lines.append(text)
            elif isinstance(block, Table):
                tables += 1
                lines.append(f"--- table {tables} ---")
                for row in block.rows:
                    cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    if any(cells):
                        lines.append(" | ".join(cells))
    except Exception as exc:
        raise DocumentLoadError(f"could not read {path.name}: {type(exc).__name__}: {exc}") from exc
    if not lines:
        lines = ["(no text in this document)"]
    return _bounded(path, lines, kind=FormatKind.WORD.value)


def load_xlsx(path: Path) -> LoadedDocument:
    """Each sheet a page, one row per line, values as the sheet shows them (not formulas)."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise DocumentLoadError(f"openpyxl is required to read {path.name}") from exc

    try:
        workbook = load_workbook(str(path), read_only=True, data_only=True)
        lines: list[str] = []
        page_starts: list[int] = []
        notes: list[str] = []
        sheets = workbook.worksheets
        for sheet in sheets[:MAX_SHEETS]:
            page_starts.append(len(lines) + 1)
            lines.append(f"--- sheet: {sheet.title} ---")
            read = 0
            for row in sheet.iter_rows(values_only=True):
                values = ["" if v is None else str(v) for v in row]
                while values and values[-1] == "":
                    values.pop()
                if not values:
                    continue
                if read >= MAX_SHEET_ROWS:
                    notes.append(
                        f"sheet {sheet.title!r}: rows beyond {MAX_SHEET_ROWS} were not read"
                    )
                    break
                lines.append(",".join(values))
                read += 1
        if len(sheets) > MAX_SHEETS:
            notes.append(f"sheets {MAX_SHEETS + 1}-{len(sheets)} were not read")
        workbook.close()
    except Exception as exc:
        raise DocumentLoadError(f"could not read {path.name}: {type(exc).__name__}: {exc}") from exc
    return _bounded(
        path,
        lines,
        kind=FormatKind.SPREADSHEET.value,
        page_starts=page_starts,
        page_count=len(page_starts),
        note="; ".join(notes),
    )


_CUE_TIME = re.compile(r"(\d{1,2}:)?\d{1,2}:\d{2}[.,]\d{1,3}\s*-->")


def load_subtitles(path: Path) -> LoadedDocument:
    """SRT or WebVTT: one cue per line, prefixed with its start time, so a citation is a moment."""
    raw = path.read_text(encoding="utf-8", errors="replace")
    lines: list[str] = []
    start = ""
    text: list[str] = []

    def flush() -> None:
        nonlocal text
        if text:
            lines.append(f"[{start}] {' '.join(text)}" if start else " ".join(text))
        text = []

    for line in raw.splitlines():
        stripped = line.strip()
        if not stripped:
            flush()
            continue
        if stripped.startswith("WEBVTT") or (stripped.isdigit() and not text):
            continue
        if _CUE_TIME.search(stripped):
            flush()
            start = stripped.split("-->")[0].strip().replace(",", ".")
            continue
        text.append(re.sub(r"<[^>]+>", "", stripped))
    flush()
    return _bounded(path, lines or ["(no cues in this file)"], kind=FormatKind.SUBTITLES.value)


def load_image(path: Path) -> LoadedDocument:
    """Size, format and orientation (Member 2's image parser); no text until OCR."""
    try:
        from PIL import Image, ImageOps
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise DocumentLoadError(f"Pillow is required to read {path.name}") from exc
    try:
        with Image.open(path) as image:
            upright = ImageOps.exif_transpose(image) or image
            width, height = upright.size
            fmt = image.format or path.suffix.lstrip(".").upper()
            rotated = upright.size != image.size
    except Exception as exc:
        raise DocumentLoadError(f"could not read {path.name}: {type(exc).__name__}: {exc}") from exc
    lines = [
        f"--- image: {path.name} ---",
        f"{fmt}, {width} x {height} pixels" + (" (EXIF orientation applied)" if rotated else ""),
        "(no text extracted: reading text from images needs OCR, Phase 39)",
    ]
    document = _bounded(path, lines, kind=FormatKind.IMAGE.value)
    return document.model_copy(update={"has_text": False})


_SIDECARS = (".srt", ".vtt")


def load_video(path: Path) -> LoadedDocument:
    """A video's text is its transcript. A subtitle file with the same name beside it is used."""
    for suffix in _SIDECARS:
        sidecar = path.with_suffix(suffix)
        if sidecar.exists():
            transcript = load_subtitles(sidecar)
            lines = [f"--- video: {path.name}, transcript from {sidecar.name} ---", transcript.text]
            return _bounded(path, lines, kind=FormatKind.VIDEO.value)
    lines = [
        f"--- video: {path.name}, {path.stat().st_size} bytes ---",
        "(no text extracted: add a transcript as a subtitle file of the same name, "
        f"{path.stem}.srt or {path.stem}.vtt; speech-to-text is not available)",
    ]
    document = _bounded(path, lines, kind=FormatKind.VIDEO.value)
    return document.model_copy(update={"has_text": False})


_PARSERS: dict[FormatKind, Callable[[Path], LoadedDocument]] = {
    FormatKind.TEXT: load_text,
    FormatKind.TABLE: load_text,
    FormatKind.SPREADSHEET: load_xlsx,
    FormatKind.WORD: load_docx,
    FormatKind.PDF: load_pdf,
    FormatKind.SUBTITLES: load_subtitles,
    FormatKind.IMAGE: load_image,
    FormatKind.VIDEO: load_video,
}
