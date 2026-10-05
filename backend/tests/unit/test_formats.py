"""Phase 38: every file type, one parser each, all becoming citable lines (Member 2's formats).

Files are generated in each test with the same libraries that read them, so their exact contents
are known and nothing binary is committed.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from app.core.config import get_settings
from app.tools.formats import FORMATS, FormatKind, parse, spec_for, supported_extensions
from app.tools.loader import DocumentLoadError, load_by_name, load_documents, resolve_document


def _docx(path: Path) -> Path:
    from docx import Document

    document = Document()
    document.add_paragraph("Project Aurora status report.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Item", "Amount (INR)"
    table.cell(1, 0).text, table.cell(1, 1).text = "Migration", "450000"
    document.add_paragraph("The approved baseline completion date is 30 April 2026.")
    document.save(str(path))
    return path


def _xlsx(path: Path) -> Path:
    from openpyxl import Workbook

    workbook = Workbook()
    budget = workbook.active
    assert budget is not None
    budget.title = "Budget"
    budget.append(["item", "amount", "approved"])
    budget.append(["migration", 450000, "no"])
    budget.append(["licences", "=B2/10", "yes"])
    schedule = workbook.create_sheet("Schedule")
    schedule.append(["milestone", "date"])
    schedule.append(["M4", "14 May 2026"])
    workbook.save(str(path))
    return path


def _png(path: Path) -> Path:
    from PIL import Image

    Image.new("RGB", (120, 80), "white").save(path)
    return path


SRT = """1
00:00:01,000 --> 00:00:04,000
Shipment 4821 arrived on 14 September.

2
00:01:02,500 --> 00:01:05,000
It was received by <i>Rahul Sharma</i>.
"""


# --- the registry ----------------------------------------------------------------------------


def test_every_extension_has_exactly_one_parser() -> None:
    extensions = [ext for spec in FORMATS for ext in spec.extensions]
    assert len(extensions) == len(set(extensions))
    assert {
        ".docx",
        ".xlsx",
        ".csv",
        ".pdf",
        ".png",
        ".mp4",
        ".srt",
        ".txt",
    } <= supported_extensions()


def test_the_upload_endpoint_accepts_exactly_the_registry() -> None:
    from app.api.v1.documents import ALLOWED

    assert ALLOWED == supported_extensions()


def test_an_unknown_type_is_refused_with_the_supported_list(tmp_path: Path) -> None:
    path = tmp_path / "notes.xyz"
    path.write_text("x")
    assert spec_for(path) is None
    with pytest.raises(DocumentLoadError, match="supported types are"):
        parse(path)


# --- each parser -------------------------------------------------------------------------------


def test_word_keeps_its_tables_in_order(tmp_path: Path) -> None:
    """Member 2's parser dropped tables; the figures are usually in them."""
    document = parse(_docx(tmp_path / "report.docx"))
    lines = document.text.splitlines()
    assert lines == [
        "Project Aurora status report.",
        "--- table 1 ---",
        "Item | Amount (INR)",
        "Migration | 450000",
        "The approved baseline completion date is 30 April 2026.",
    ]
    assert document.kind == "word" and document.parser == "python-docx"
    assert len(document.sha256) == 64 and document.has_text


def test_a_spreadsheet_is_one_page_per_sheet(tmp_path: Path) -> None:
    document = parse(_xlsx(tmp_path / "budget.xlsx"))
    lines = document.text.splitlines()
    assert lines[0] == "--- sheet: Budget ---"
    assert "migration,450000,no" in lines
    assert "--- sheet: Schedule ---" in lines
    assert "M4,14 May 2026" in lines
    # Each sheet starts a page, so a tool can cite `budget.xlsx:p2` for the Schedule sheet.
    assert document.page_count == 2
    assert document.page_starts == [1, lines.index("--- sheet: Schedule ---") + 1]


def test_subtitles_become_timed_lines(tmp_path: Path) -> None:
    path = tmp_path / "interview.srt"
    path.write_text(SRT, encoding="utf-8")
    document = parse(path)
    assert document.text.splitlines() == [
        "[00:00:01.000] Shipment 4821 arrived on 14 September.",
        "[00:01:02.500] It was received by Rahul Sharma.",
    ]
    vtt = tmp_path / "interview.vtt"
    vtt.write_text("WEBVTT\n\n00:00:01.000 --> 00:00:04.000\nHello there.\n", encoding="utf-8")
    assert parse(vtt).text == "[00:00:01.000] Hello there."


def test_an_image_is_described_and_says_it_has_no_text_yet(tmp_path: Path) -> None:
    document = parse(_png(tmp_path / "scan.png"))
    assert document.kind == "image" and document.parser == "pillow"
    assert "120 x 80 pixels" in document.text
    assert not document.has_text
    # Understanding (OCR, the vision model) is off in unit tests; the line says so.
    assert "MEDIA_UNDERSTANDING" in document.text.splitlines()[-1]


def test_a_video_reads_its_transcript_when_one_sits_beside_it(tmp_path: Path) -> None:
    video = tmp_path / "meeting.mp4"
    video.write_bytes(b"\x00\x00\x00\x18ftypmp42")
    alone = parse(video)
    assert not alone.has_text
    assert "meeting.srt" in alone.text

    (tmp_path / "meeting.srt").write_text(SRT, encoding="utf-8")
    with_transcript = parse(video)
    assert with_transcript.has_text
    assert "[00:00:01.000] Shipment 4821 arrived on 14 September." in with_transcript.text
    assert with_transcript.kind == FormatKind.VIDEO.value


def test_text_and_csv_read_as_before(tmp_path: Path) -> None:
    csv = tmp_path / "budget.csv"
    csv.write_text("item,amount\nmigration,450000\n", encoding="utf-8")
    document = parse(csv)
    assert document.kind == "csv" and document.parser == "text"
    assert document.text.startswith("item,amount")


def test_a_corrupt_office_file_is_an_error_not_an_empty_document(tmp_path: Path) -> None:
    broken = tmp_path / "broken.docx"
    broken.write_bytes(b"not a zip")
    with pytest.raises(DocumentLoadError, match=r"could not read broken\.docx"):
        parse(broken)


# --- what a mission reads --------------------------------------------------------------------


def test_a_mission_excludes_files_with_no_text_yet(tmp_path: Path) -> None:
    text = tmp_path / "notes.txt"
    text.write_text("Shipment 4821 arrived on 14 September.\n", encoding="utf-8")
    documents, loaded = load_documents([text, _png(tmp_path / "photo.png")])
    assert list(documents) == ["notes.txt"]
    assert [d.document_id for d in loaded] == ["notes.txt"]


def test_an_uploaded_file_can_be_used_by_name(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """BUG-023: uploads land in .agent/uploads/, and no mission could find them by name."""
    # `agent_dir` is a property; point it at a temporary `.agent` for this test only.
    monkeypatch.setattr(type(get_settings()), "agent_dir", property(lambda _self: tmp_path))
    uploads = tmp_path / "uploads"
    uploads.mkdir()
    _docx(uploads / "field_report.docx")

    assert resolve_document("field_report.docx") == uploads / "field_report.docx"
    documents, _ = load_by_name(["field_report.docx"])
    assert "Migration | 450000" in documents["field_report.docx"]
