"""Phase 51: the investigation report - one object, three exports, every export on the record.

The executive summary and the risk come only from findings that count; severity is set by declared
rules; the currency check catches a claim whose unit is not its evidence's; the PDF, DOCX and
Markdown are drawn from the same sections and each carries the content hash; and exporting writes a
REPORTED event into the custody of every cited document and an entry into the audit log.
"""

from __future__ import annotations

import html
import io
from pathlib import Path

import pytest
from docx import Document
from httpx import ASGITransport, AsyncClient
from pypdf import PdfReader

from app.api.app import create_app
from app.api.registry import MissionRecord, get_registry
from app.evidence import service as evidence_service
from app.orchestration.mission import MissionResult, MissionStatus, Stage
from app.reporting.builder import build_report, content_hash, get_severity_rules, severity_of
from app.reporting.document import build_document
from app.reporting.export import to_markdown
from app.reporting.render import render_docx, render_html, render_pdf
from app.reporting.units import currencies, unit_mismatch
from app.review.service import get_review_store
from app.schemas.common import SourceLocator, new_run_id
from app.schemas.evidence import EvidenceRef, ResolutionStatus
from app.schemas.evidence_locker import EvidenceItem
from app.schemas.finding import Finding, StatementType
from app.schemas.investigation_report import Severity
from app.schemas.objective import AttachedDocument, Objective, ObjectiveScope
from app.schemas.review import FindingReview, ReviewDecision, ReviewState
from app.schemas.verification import VerificationResult, VerificationStatus


@pytest.fixture(autouse=True)
def _versions_in_tmp(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    """Report versions are written beside the run archives; tests keep them out of `.agent/`."""
    from app.reporting import versions

    monkeypatch.setattr(versions, "_path", lambda run_id: tmp_path / f"{run_id}.json")


def _finding(
    finding_id: str,
    claim: str,
    cite: str,
    *,
    statement: StatementType = StatementType.SOURCE_CLAIM,
    verified: bool = True,
) -> Finding:
    document, row = cite.rsplit(":r", 1)
    return Finding(
        finding_id=finding_id,
        claim=claim,
        evidence=[
            EvidenceRef(
                evidence_id="E-001",
                locator=SourceLocator(document_id=document, document_name=document, row=int(row)),
                resolution=ResolutionStatus.RESOLVED,
            )
        ],
        statement_type=statement,
        statement_basis=StatementType.SOURCE_CLAIM,
        verification=VerificationResult(status=VerificationStatus.SUPPORTED, confidence=0.9)
        if verified
        else None,
    )


TEXT = {
    "ledger.txt:r3": "Spend to date: INR 287,500 against the approved INR 250,000.",
    "plan.txt:r2": "The go-live date is 12 June 2026.",
}


# --- the currency check ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("text", "found"),
    [
        ("$287,500", {"USD"}),
        ("INR 287,500", {"INR"}),
        ("Rs. 2.5 lakh", {"INR"}),
        ("₹ 40,000", {"INR"}),
        ("EUR 10,000 and £5", {"EUR", "GBP"}),
        ("287,500", set()),
    ],
)
def test_currencies_are_read_from_symbols_codes_and_words(text: str, found: set[str]) -> None:
    assert currencies(text) == found


def test_a_dollar_claim_over_rupee_evidence_is_a_unit_mismatch() -> None:
    assert unit_mismatch("Spend reached $287,500.", [TEXT["ledger.txt:r3"]]) == (
        "the claim states USD where its cited lines state INR"
    )
    assert unit_mismatch("Spend reached INR 287,500.", [TEXT["ledger.txt:r3"]]) == ""
    assert unit_mismatch("Spend reached 287,500.", [TEXT["ledger.txt:r3"]]) == ""


# --- severity -------------------------------------------------------------------------------------


def test_severity_is_the_highest_rule_and_says_which() -> None:
    rules = get_severity_rules()
    plain = _finding("F-001", "x", "plan.txt:r2")
    assert severity_of(plain, unit_issue="", challenge=None, case_severity="", rules=rules) == (
        Severity.LOW,
        "no rule raised it",
    )
    contradiction = _finding("F-002", "x", "plan.txt:r2", statement=StatementType.CONTRADICTION)
    assert severity_of(
        contradiction, unit_issue="", challenge=None, case_severity="", rules=rules
    ) == (Severity.HIGH, "it reports a contradiction")
    assert severity_of(plain, unit_issue="", challenge=None, case_severity="HIGH", rules=rules) == (
        Severity.MEDIUM,
        "its case is HIGH",
    )
    assert (
        severity_of(plain, unit_issue="USD vs INR", challenge=None, case_severity="", rules=rules)[
            0
        ]
        is Severity.HIGH
    )


# --- the report -----------------------------------------------------------------------------------


def _result() -> MissionResult:
    return MissionResult(
        run_id=new_run_id(),
        objective=Objective(
            text="Did Orion overspend?",
            scope=ObjectiveScope(
                documents=[
                    AttachedDocument(document_id="ledger.txt", name="ledger.txt"),
                    AttachedDocument(document_id="plan.txt", name="plan.txt"),
                ]
            ),
        ),
        findings=[
            _finding("F-001", "The go-live date is 12 June 2026.", "plan.txt:r2"),
            _finding(
                "F-002",
                "Spend of INR 287,500 exceeds the approved INR 250,000.",
                "ledger.txt:r3",
                statement=StatementType.CONTRADICTION,
            ),
            _finding("F-003", "Spend reached $287,500.", "ledger.txt:r3"),
            _finding("F-004", "The vendor was paid twice.", "ledger.txt:r3", verified=False),
        ],
        evidence_text=TEXT,
    )


def _pending(finding_id: str) -> FindingReview:
    return FindingReview(
        run_id="r",
        finding_id=finding_id,
        claim="",
        required=True,
        required_because=["it reports a contradiction"],
    )


def test_only_findings_that_count_reach_the_summary_and_the_risk() -> None:
    report = build_report(_result(), reviews={"F-002": _pending("F-002")})

    assert [f.finding_id for f in report.established] == ["F-001"]
    assert [f.finding_id for f in report.awaiting_review] == ["F-002"]
    assert {f.finding_id for f in report.not_established} == {"F-003", "F-004"}
    assert report.executive_summary == ["The go-live date is 12 June 2026. [F-001; plan.txt:r2]"]
    assert report.risk is Severity.LOW, "the HIGH contradiction awaits review and cannot raise it"
    unit = next(f for f in report.not_established if f.finding_id == "F-003")
    assert unit.unit_issues == ["the claim states USD where its cited lines state INR"]
    assert unit.severity is Severity.HIGH


def test_approval_moves_a_finding_into_the_summary_and_the_risk() -> None:
    approved = _pending("F-002").model_copy(
        update={
            "state": ReviewState.APPROVED,
            "history": [
                ReviewDecision(
                    run_id="r", finding_id="F-002", state=ReviewState.APPROVED, reviewer="lead"
                )
            ],
        }
    )
    report = build_report(_result(), reviews={"F-002": approved})
    assert {f.finding_id for f in report.established} == {"F-001", "F-002"}
    assert report.risk is Severity.HIGH and report.risk_reason.startswith("F-002")


def test_the_hash_covers_the_content() -> None:
    report = build_report(_result())
    assert report.content_hash == content_hash(report)
    changed = report.model_copy(update={"objective": "Something else?"})
    assert content_hash(changed) != report.content_hash
    rebuilt = build_report(report_result := _result())
    assert build_report(report_result).content_hash == rebuilt.content_hash, "time is not content"


def test_the_timeline_and_recommendations_come_from_the_run() -> None:
    report = build_report(_result(), reviews={"F-002": _pending("F-002")})
    assert report.timeline[0].date == "2026-06-12"
    assert "Decide on F-002" in report.recommendations[0]
    assert any("Establish which source is current for F-002" in r for r in report.recommendations)


# --- the exports ----------------------------------------------------------------------------------


def test_every_export_carries_the_same_content_and_the_hash() -> None:
    """v4 Phase 5: screen, PDF and DOCX are drawn from one document, so they hold the same words."""
    report = build_report(_result(), receipt_hash="r" * 64)
    document = build_document(report)
    claims = [
        f.claim for f in [*report.established, *report.awaiting_review, *report.not_established]
    ]
    assert claims, "the fixture has findings to show"

    markdown = to_markdown(report)
    assert "## Executive summary" in markdown and report.content_hash in markdown

    page = render_html(document)
    assert page.startswith("<!doctype html>") and "<script" not in page
    for words in ("Executive brief", "Findings", "Appendix B - Evidence and file hashes"):
        assert words in page
    assert report.content_hash[:16] in page and report.receipt_hash[:16] in page

    pdf = PdfReader(io.BytesIO(render_pdf(document)))
    text = " ".join(" ".join(p.extract_text().split()) for p in pdf.pages)
    assert "Investigation Report" in text and "Executive brief" in text and "Findings" in text
    assert report.content_hash[:16] in text and "Page 1" in text

    docx = Document(io.BytesIO(render_docx(document)))
    headings = [p.text for p in docx.paragraphs if p.style.name.startswith("Heading")]
    assert headings[:2] == ["Executive brief", "Findings"]
    assert "Appendix A - Receipt" in headings
    cells = {c.text for t in docx.tables for row in t.rows for c in row.cells}
    assert report.receipt_hash in cells
    assert report.content_hash[:16] in docx.sections[0].footer.paragraphs[0].text

    # The same findings on screen, in the PDF and in the DOCX.
    for claim in claims:
        assert html.escape(claim) in page
        assert " ".join(claim.split())[:40] in text
        assert claim in cells


def test_a_contradiction_is_shown_with_both_sides_quoted() -> None:
    from app.schemas.investigation_report import ReportFinding, Severity

    report = build_report(_result())
    report.established = [
        ReportFinding(
            finding_id="F-009",
            claim="The completion dates differ",
            statement="CONTRADICTION",
            severity=Severity.MEDIUM,
            severity_reason="dates",
            confidence=0.8,
            confidence_explained="",
            verification="SUPPORTED",
            review="not required",
            citations=["plan.txt:r2", "ledger.txt:r5"],
            quotes={
                "plan.txt:r2": "Completion: 30 April 2026",
                "ledger.txt:r5": "Closed 14 May 2026",
            },
        )
    ]
    pair = next(b for b in build_document(report).blocks if b.kind == "pair")
    assert pair.left == ("plan.txt:r2", "Completion: 30 April 2026")
    assert pair.right == ("ledger.txt:r5", "Closed 14 May 2026")
    assert '"Closed 14 May 2026"' in render_html(build_document(report))


def test_the_rupee_sign_prints_or_is_spelled_out() -> None:
    report = build_report(_result())
    report.objective = "Spend → ₹ 40,000 — “approved”"
    text = PdfReader(io.BytesIO(render_pdf(build_document(report)))).pages[0].extract_text()
    assert "₹ 40,000" in text or "INR  40,000" in text


# --- the API --------------------------------------------------------------------------------------


async def test_an_export_is_written_to_custody_and_the_audit_log() -> None:
    result = _result()
    record = MissionRecord(result.run_id, result.objective, ["ledger.txt", "plan.txt"])
    record.status = MissionStatus.COMPLETED
    record.stage = Stage.DONE
    record.result = result
    locker = await evidence_service.get_locker()
    item = await locker.add(EvidenceItem(evidence_id="EV-000901", sha256="a" * 64, name="plan.txt"))
    record.evidence = {"plan.txt": item.evidence_id}
    get_registry()._missions[record.run_id] = record

    async with AsyncClient(transport=ASGITransport(app=create_app()), base_url="http://test") as c:
        report = (await c.get(f"/api/v1/missions/{record.run_id}/investigation-report")).json()
        pdf = await c.get(
            f"/api/v1/missions/{record.run_id}/investigation-report/export/pdf",
            headers={"X-Jarvis-User": "lead-1"},
        )
        bad = await c.get(f"/api/v1/missions/{record.run_id}/investigation-report/export/xls")

    assert report["evidence"][1]["evidence_id"] == "EV-000901"
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF-1.4")
    assert pdf.headers["X-Report-Hash"] == report["content_hash"]
    assert bad.status_code == 422

    custody = [e for e in await locker.history("EV-000901") if e.action.value == "REPORTED"]
    assert custody and custody[0].actor == "lead-1"
    assert custody[0].details["content_hash"] == report["content_hash"]
    audit = await (await get_review_store()).audit(target=record.run_id)
    assert any(e.action == "report.exported" and e.details["format"] == "pdf" for e in audit)
    # The HIGH contradiction was challenged before it was reported (Phase 48's deferred step).
    assert any(e.action == "challenge.before_report" for e in audit)
    assert "F-002" in record.challenges and "F-001" not in record.challenges


async def test_an_unfinished_run_has_no_report() -> None:
    record = MissionRecord(new_run_id(), Objective(text="Still running?"), [])
    record.result = MissionResult(run_id=record.run_id, objective=record.objective)
    get_registry()._missions[record.run_id] = record
    async with AsyncClient(transport=ASGITransport(app=create_app()), base_url="http://test") as c:
        response = await c.get(f"/api/v1/missions/{record.run_id}/investigation-report")
    assert response.status_code == 409


async def test_each_export_is_a_signed_version_of_the_report() -> None:
    """v4 Phase 5: every export is numbered and signed with the content, receipt and case-snapshot
    hashes; the HTML opens in the browser, the files download."""
    result = _result()
    record = MissionRecord(result.run_id, result.objective, ["ledger.txt", "plan.txt"])
    record.status = MissionStatus.COMPLETED
    record.stage = Stage.DONE
    record.result = result
    get_registry()._missions[record.run_id] = record
    base = f"/api/v1/missions/{record.run_id}/investigation-report"
    async with AsyncClient(transport=ASGITransport(app=create_app()), base_url="http://test") as c:
        page = await c.get(f"{base}/export/html", headers={"X-Jarvis-User": "lead-1"})
        docx = await c.get(f"{base}/export/docx")
        versions = (await c.get(f"{base}/versions")).json()

    assert page.status_code == 200 and page.headers["content-type"].startswith("text/html")
    assert page.headers["content-disposition"].startswith("inline")
    assert docx.headers["content-disposition"].startswith("attachment")
    assert docx.headers["X-Report-Version"] == "2"
    assert [(v["version"], v["format"]) for v in versions] == [(1, "html"), (2, "docx")]
    first, second = versions
    assert first["actor"] == "lead-1"
    assert first["receipt_hash"] and first["receipt_hash"] == page.headers["X-Receipt-Hash"]
    # Nothing changed between the two exports: the same state of the case, the same snapshot.
    assert first["snapshot_hash"] == second["snapshot_hash"]
    assert first["findings"] and first["findings"] == second["findings"]
    assert set(first["findings"]) == {f.finding_id for f in result.findings}
